from __future__ import annotations

import json
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from enso_scenario import SOURCE_NAME as ENSO_SOURCE_NAME
from enso_scenario import anomaly_for_date, save_scenario_to_database

from config import (
    ANALOG_TOP_N,
    BASE_DIR,
    BASIN_NAMES,
    END_DATE,
    FORECAST_FILE,
    HOCHUA_DB_PATH,
    MAX_DAILY_WATER_CHANGE,
    MODEL_DIR,
    NSMO_WEATHER_DB_PATH,
    OUTPUT_DIR,
    RESERVOIR_OPERATION_RULES,
    WATER_LEVEL_BOUNDS,
    WEATHER_IMPACT_BASIN_WEIGHTS,
    WEATHER_IMPACT_DAYS,
    WEATHER_RAIN_KEYWORDS,
    ensure_project_dirs,
)

ANALOG_YEAR = 2025
FORECAST_START_DATE = "2026-01-01"
FORECAST_END_DATE = "2026-12-31"
FORECAST_YEAR = 2026
STORM_MONTHS = {9, 10, 11}
RAINY_SHOULDER_MONTHS = {8, 12}
KTTV_DIR = BASE_DIR / "KTTV"
KTTV_SOURCE_WEIGHTS = {
    "Truot5Ngay": 0.78,
    "Tuan": 0.55,
    "Thang": 0.35,
}
KTTV_WATER_RESPONSE = {
    "BTS": 0.08,
    "BK": 0.05,
    "SP3": 0.05,
}
KTTV_FORECAST_CACHE: pd.DataFrame | None = None
RESERVOIR_BY_BASIN = {
    "BTS": "Buôn Tua Srah",
    "BK": "Buôn Kuốp",
    "SP3": "Srêpốk 3",
}


def normalize_text(value: object) -> str:
    text = "" if pd.isna(value) else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"\s+", " ", text.lower()).strip()


def load_current_weather_impact() -> dict[str, object]:
    """Read latest NSMO weather rows and convert them to a qualitative rain impact score."""
    if not NSMO_WEATHER_DB_PATH.exists():
        return {"date": "", "regions": {}, "wetness": 0.0, "summary": "no_weather_db"}
    try:
        with sqlite3.connect(NSMO_WEATHER_DB_PATH) as conn:
            latest_row = conn.execute("SELECT MAX(report_date) FROM nsmo_weather_region").fetchone()
            latest_date = latest_row[0] if latest_row else None
            if not latest_date:
                return {"date": "", "regions": {}, "wetness": 0.0, "summary": "no_weather_rows"}
            rows = pd.read_sql_query(
                """
                SELECT region_key, city, temp_c, humidity, condition_text
                FROM nsmo_weather_region
                WHERE report_date = ?
                """,
                conn,
                params=(latest_date,),
            )
    except Exception as exc:
        return {"date": "", "regions": {}, "wetness": 0.0, "summary": f"weather_error:{exc}"}

    regions: dict[str, dict[str, object]] = {}
    wetness_scores = []
    for row in rows.to_dict("records"):
        condition = normalize_text(row.get("condition_text"))
        humidity = pd.to_numeric(row.get("humidity"), errors="coerce")
        keyword_hit = any(keyword in condition for keyword in [normalize_text(k) for k in WEATHER_RAIN_KEYWORDS])
        score = 0.0
        if keyword_hit:
            score += 0.55
        if pd.notna(humidity):
            score += max(0.0, min(0.35, (float(humidity) - 65.0) / 100.0))
        score = float(min(score, 1.0))
        region_key = str(row.get("region_key") or "").strip()
        regions[region_key] = {
            "city": row.get("city") or "",
            "condition": row.get("condition_text") or "",
            "humidity": None if pd.isna(humidity) else float(humidity),
            "wetness": score,
        }
        wetness_scores.append(score)
    return {
        "date": latest_date,
        "regions": regions,
        "wetness": float(np.mean(wetness_scores)) if wetness_scores else 0.0,
        "summary": "; ".join(
            f"{value['city']}: {value['condition']}" for value in regions.values() if value.get("condition")
        ),
    }


def basin_weather_wetness(basin_id: str, weather: dict[str, object]) -> float:
    regions = weather.get("regions") if isinstance(weather, dict) else {}
    if not isinstance(regions, dict):
        return 0.0
    weights = WEATHER_IMPACT_BASIN_WEIGHTS.get(basin_id, {})
    if not weights:
        return float(weather.get("wetness") or 0.0)
    total = 0.0
    weight_sum = 0.0
    for region, weight in weights.items():
        data = regions.get(region) or {}
        total += float(weight) * float(data.get("wetness") or 0.0)
        weight_sum += float(weight)
    return float(total / weight_sum) if weight_sum else 0.0


def parse_kttv_basin(path: Path) -> str | None:
    name = normalize_text(path.name)
    if "buontuasrah" in name or "buon tua srah" in name:
        return "BTS"
    if "buonkuop" in name or "buon kuop" in name:
        return "BK"
    if "srepok" in name or "sre pok" in name:
        return "SP3"
    return None


def parse_kttv_source_type(path: Path) -> str:
    name = normalize_text(path.name)
    if "truot5ngay" in name:
        return "Truot5Ngay"
    if "tuan" in name:
        return "Tuan"
    if "thang" in name:
        return "Thang"
    return "Khac"


def parse_kttv_issue_date(path: Path) -> pd.Timestamp | None:
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", path.name)
    if not match:
        return None
    day, month, year = [int(part) for part in match.groups()]
    return pd.Timestamp(year=year, month=month, day=day)


def to_kttv_date(value: object) -> pd.Timestamp | None:
    if value is None or pd.isna(value):
        return None
    parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return pd.Timestamp(parsed).normalize()


def to_float(value: object) -> float | None:
    number = pd.to_numeric(value, errors="coerce")
    if pd.isna(number):
        return None
    return float(number)


def parse_month_year_from_sheet(rows: list[tuple[object, ...]]) -> tuple[int, int] | None:
    pattern = re.compile(r"thang\s+(\d{1,2})\s+nam\s+(\d{4})")
    for row in rows:
        for value in row:
            text = normalize_text(value)
            match = pattern.search(text)
            if match:
                month, year = [int(part) for part in match.groups()]
                return month, year
    return None


def expand_date_range(start: pd.Timestamp, end: pd.Timestamp) -> list[pd.Timestamp]:
    if start > end:
        return []
    return list(pd.date_range(start, end, freq="D"))


def parse_kttv_file(path: Path) -> list[dict[str, object]]:
    basin_id = parse_kttv_basin(path)
    source_type = parse_kttv_source_type(path)
    issue_date = parse_kttv_issue_date(path)
    if not basin_id or source_type == "Khac":
        return []

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []

    rows = [tuple(row) for worksheet in workbook.worksheets for row in worksheet.iter_rows(values_only=True)]
    records: list[dict[str, object]] = []
    source_priority = {"Truot5Ngay": 3, "Tuan": 2, "Thang": 1}.get(source_type, 0)

    if source_type == "Truot5Ngay":
        q_col = 2 if basin_id == "BTS" else 3
        for row in rows:
            if len(row) <= q_col:
                continue
            date = to_kttv_date(row[0])
            q_value = to_float(row[q_col])
            if date is None or q_value is None:
                continue
            time_text = normalize_text(row[1] if len(row) > 1 else "")
            if time_text != "ngay":
                continue
            records.append(
                {
                    "Date": date,
                    "Basin_ID": basin_id,
                    "KTTV_Q_m3s": q_value,
                    "Source_Type": source_type,
                    "Issue_Date": issue_date,
                    "Source_File": path.name,
                    "Source_Priority": source_priority,
                }
            )

    elif source_type == "Tuan":
        q_col = 4 if basin_id == "BTS" else 5
        for row in rows:
            if len(row) <= q_col:
                continue
            period = normalize_text(row[0])
            characteristic = normalize_text(row[3] if len(row) > 3 else "")
            if not period.startswith("w") or characteristic != "qtb":
                continue
            start = to_kttv_date(row[1])
            end = to_kttv_date(row[2])
            q_value = to_float(row[q_col])
            if start is None or end is None or q_value is None:
                continue
            for date in expand_date_range(start, end):
                records.append(
                    {
                        "Date": date,
                        "Basin_ID": basin_id,
                        "KTTV_Q_m3s": q_value,
                        "Source_Type": source_type,
                        "Issue_Date": issue_date,
                        "Source_File": path.name,
                        "Source_Priority": source_priority,
                    }
                )

    elif source_type == "Thang":
        month_year = parse_month_year_from_sheet(rows)
        q_col = 2 if basin_id == "BTS" else 3
        if month_year:
            month, year = month_year
            month_start = pd.Timestamp(year=year, month=month, day=1)
            month_end = month_start + pd.offsets.MonthEnd(0)
            for row in rows:
                if len(row) <= q_col:
                    continue
                characteristic = normalize_text(row[0])
                forecast_flag = normalize_text(row[1] if len(row) > 1 else "")
                q_value = to_float(row[q_col])
                if "qve tb thang" not in characteristic or "du bao" not in forecast_flag or q_value is None:
                    continue
                for date in expand_date_range(month_start, month_end):
                    records.append(
                        {
                            "Date": date,
                            "Basin_ID": basin_id,
                            "KTTV_Q_m3s": q_value,
                            "Source_Type": source_type,
                            "Issue_Date": issue_date,
                            "Source_File": path.name,
                            "Source_Priority": source_priority,
                        }
                    )
                break
    return records


def load_kttv_forecast_inputs() -> pd.DataFrame:
    if not KTTV_DIR.exists():
        return pd.DataFrame()
    records: list[dict[str, object]] = []
    for path in sorted(KTTV_DIR.glob("*.xlsx")):
        records.extend(parse_kttv_file(path))
    if not records:
        return pd.DataFrame()

    df = pd.DataFrame(records)
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Issue_Date"] = pd.to_datetime(df["Issue_Date"], errors="coerce")
    df["KTTV_Q_m3s"] = pd.to_numeric(df["KTTV_Q_m3s"], errors="coerce")
    df = df.dropna(subset=["Date", "Basin_ID", "KTTV_Q_m3s"])
    df = df.sort_values(["Basin_ID", "Date", "Source_Priority", "Issue_Date"], ascending=[True, True, False, False])
    return df.drop_duplicates(["Basin_ID", "Date"], keep="first").reset_index(drop=True)


def refresh_kttv_forecast_database() -> pd.DataFrame:
    global KTTV_FORECAST_CACHE
    df = load_kttv_forecast_inputs()
    KTTV_FORECAST_CACHE = df
    if df.empty:
        return df
    sqlite_df = df.copy()
    sqlite_df["Date"] = sqlite_df["Date"].dt.strftime("%Y-%m-%d")
    sqlite_df["Issue_Date"] = sqlite_df["Issue_Date"].dt.strftime("%Y-%m-%d")
    with sqlite3.connect(HOCHUA_DB_PATH) as conn:
        sqlite_df.to_sql("kttv_forecast_daily", conn, if_exists="replace", index=False)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_kttv_forecast_daily_basin_date ON kttv_forecast_daily(Basin_ID, Date)")
    print(f"Da nap du bao KTTV: {len(sqlite_df)} dong tu {KTTV_DIR}")
    return df


def kttv_series_for_basin(basin_id: str, dates: pd.DatetimeIndex) -> pd.DataFrame:
    global KTTV_FORECAST_CACHE
    if KTTV_FORECAST_CACHE is None:
        KTTV_FORECAST_CACHE = load_kttv_forecast_inputs()
    if KTTV_FORECAST_CACHE.empty:
        return pd.DataFrame()
    df = KTTV_FORECAST_CACHE[KTTV_FORECAST_CACHE["Basin_ID"].eq(basin_id)].copy()
    if df.empty:
        return df
    date_set = set(pd.to_datetime(dates).normalize())
    df = df[df["Date"].isin(date_set)].copy()
    return df.sort_values("Date")


def load_dashboard_daily_actuals(basin_id: str) -> pd.DataFrame:
    reservoir_name = RESERVOIR_BY_BASIN.get(basin_id)
    if not reservoir_name or not HOCHUA_DB_PATH.exists():
        return pd.DataFrame(columns=["Date", "Basin_ID", "Q_in_m3s", "WaterLevel_m"])
    try:
        with sqlite3.connect(HOCHUA_DB_PATH) as conn:
            df = pd.read_sql_query(
                """
                SELECT
                    date(observed_at) AS Date,
                    AVG(q_ve) AS Q_in_m3s,
                    AVG(h_tl) AS WaterLevel_m
                FROM (
                    SELECT *
                    FROM (
                        SELECT
                            *,
                            ROW_NUMBER() OVER (
                                PARTITION BY observed_at
                                ORDER BY requested_at_last_seen DESC, id DESC
                            ) AS rn
                        FROM reservoir_readings
                        WHERE reservoir_name LIKE ?
                          AND observed_at IS NOT NULL
                          AND (q_ve IS NOT NULL OR h_tl IS NOT NULL)
                    )
                    WHERE rn = 1
                )
                GROUP BY date(observed_at)
                ORDER BY Date
                """,
                conn,
                params=(f"{reservoir_name}%",),
                parse_dates=["Date"],
            )
    except Exception as exc:
        print(f"Bo qua du lieu thuc do dashboard {basin_id}: {exc}")
        return pd.DataFrame(columns=["Date", "Basin_ID", "Q_in_m3s", "WaterLevel_m"])
    if df.empty:
        return df
    df["Basin_ID"] = basin_id
    for column in ["Q_in_m3s", "WaterLevel_m"]:
        df[column] = pd.to_numeric(df[column], errors="coerce")
    low, high = WATER_LEVEL_BOUNDS[basin_id]
    df.loc[~df["WaterLevel_m"].between(low, high), "WaterLevel_m"] = np.nan
    return df[["Date", "Basin_ID", "Q_in_m3s", "WaterLevel_m"]]


def refresh_recent_hydro_features(history: pd.DataFrame, basin_id: str) -> pd.DataFrame:
    history = history.sort_values("Date").copy()
    if "Q_in_m3s" in history.columns:
        q = pd.to_numeric(history["Q_in_m3s"], errors="coerce")
        q_history = q.shift(1)
        for lag in [1, 2, 3, 7, 14, 30]:
            history[f"Q_Lag{lag}"] = q.shift(lag)
        for window in [3, 7, 15, 30]:
            history[f"Q_{window}d_mean"] = q_history.rolling(window, min_periods=1).mean()
            history[f"Q_{window}d_max"] = q_history.rolling(window, min_periods=1).max()
    if "WaterLevel_m" in history.columns:
        water = pd.to_numeric(history["WaterLevel_m"], errors="coerce")
        low, high = WATER_LEVEL_BOUNDS[basin_id]
        water = water.where(water.between(low, high))
        history["WaterLevel_m"] = water
        water_history = water.shift(1)
        for lag in [1, 2, 3, 7, 14, 30]:
            history[f"WaterLevel_Lag{lag}"] = water.shift(lag)
        for window in [3, 7, 15, 30]:
            history[f"WaterLevel_{window}d_mean"] = water_history.rolling(window, min_periods=1).mean()
            history[f"WaterLevel_{window}d_min"] = water_history.rolling(window, min_periods=1).min()
            history[f"WaterLevel_{window}d_max"] = water_history.rolling(window, min_periods=1).max()
    return history


def augment_history_with_dashboard_actuals(history: pd.DataFrame, basin_id: str) -> pd.DataFrame:
    actuals = load_dashboard_daily_actuals(basin_id)
    if actuals.empty:
        return refresh_recent_hydro_features(history, basin_id)
    history = history.copy()
    history["Date"] = pd.to_datetime(history["Date"], errors="coerce").dt.normalize()
    history = history.dropna(subset=["Date"]).sort_values("Date")
    history = history.set_index("Date")
    actuals = actuals.set_index("Date")
    for date, row in actuals.iterrows():
        if date not in history.index:
            history.loc[date, "Basin_ID"] = basin_id
        for column in ["Q_in_m3s", "WaterLevel_m"]:
            value = row.get(column)
            if pd.notna(value):
                history.loc[date, column] = value
    history["Basin_ID"] = history.get("Basin_ID", basin_id).fillna(basin_id)
    history = history.reset_index().sort_values("Date")
    return refresh_recent_hydro_features(history, basin_id)


def read_rmse(basin_id: str, slug: str, fallback: float) -> float:
    candidates = sorted(
        MODEL_DIR.glob(f"metrics_{basin_id}_{slug}*.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if slug == "Q":
        candidates.extend(
            sorted(
                MODEL_DIR.glob(f"metrics_{basin_id}*.json"),
                key=lambda path: path.stat().st_mtime,
                reverse=True,
            )
        )
    metric_file = candidates[0] if candidates else MODEL_DIR / f"metrics_{basin_id}_{slug}.json"
    if not metric_file.exists():
        return fallback
    with open(metric_file, encoding="utf-8") as file:
        return float(json.load(file).get("RMSE", fallback))


def read_conformal_residuals(basin_id: str, slug: str) -> pd.DataFrame:
    """Read backtest residuals used for month-specific conformal intervals."""
    target = "Q_in_m3s" if slug == "Q" else "WaterLevel_m"
    pred_col = f"{slug}_Predicted"
    path = OUTPUT_DIR / f"Backtest_{basin_id}_{slug}.csv"
    if not path.exists():
        return pd.DataFrame(columns=["Month", "Residual"])
    df = pd.read_csv(path, parse_dates=["Date"])
    if target not in df.columns or pred_col not in df.columns:
        return pd.DataFrame(columns=["Month", "Residual"])
    residuals = pd.DataFrame(
        {
            "Month": df["Date"].dt.month,
            "Residual": pd.to_numeric(df[pred_col], errors="coerce") - pd.to_numeric(df[target], errors="coerce"),
        }
    ).dropna()
    return residuals


def conformal_bounds(center: pd.Series, residuals: pd.DataFrame, floor: float | None = None) -> tuple[pd.Series, pd.Series]:
    """Create dynamic P10/P90 bounds from empirical residuals by month."""
    lower = pd.Series(np.nan, index=center.index, dtype="float64")
    upper = pd.Series(np.nan, index=center.index, dtype="float64")
    if residuals.empty:
        spread = float(center.std(skipna=True) or 0.0) * 0.25
        lower = center - spread
        upper = center + spread
    else:
        global_q10 = float(residuals["Residual"].quantile(0.10))
        global_q90 = float(residuals["Residual"].quantile(0.90))
        for date in center.index:
            month_values = residuals.loc[residuals["Month"] == date.month, "Residual"]
            q10 = float(month_values.quantile(0.10)) if len(month_values) >= 20 else global_q10
            q90 = float(month_values.quantile(0.90)) if len(month_values) >= 20 else global_q90
            lower.loc[date] = center.loc[date] + q10
            upper.loc[date] = center.loc[date] + q90
    if floor is not None:
        lower = lower.clip(lower=floor)
        upper = upper.clip(lower=floor)
    return lower, upper


def same_month_day(history: pd.DataFrame, target_date: pd.Timestamp, column: str) -> float | None:
    month_day = target_date.strftime("%m-%d")
    climatology = history[
        (history["Date"].dt.month == target_date.month)
        & (history["Date"].dt.day == target_date.day)
    ]
    values = pd.to_numeric(climatology[column], errors="coerce").dropna()
    if not values.empty:
        return float(values.median())

    day_gap = (history["Date"].dt.dayofyear - target_date.dayofyear).abs()
    nearby = history[day_gap <= 15]
    values = pd.to_numeric(nearby[column], errors="coerce").dropna()
    if not values.empty:
        return float(values.median())

    monthly = history[history["Date"].dt.month == target_date.month]
    values = pd.to_numeric(monthly[column], errors="coerce").dropna()
    if not values.empty:
        return float(values.median())

    values = pd.to_numeric(history[column], errors="coerce").dropna()
    return float(values.median()) if not values.empty else None


def historical_value_for_date(history: pd.DataFrame, year: int, target_date: pd.Timestamp, column: str) -> float | None:
    try:
        analog_date = target_date.replace(year=year)
    except ValueError:
        analog_date = target_date.replace(year=year, day=28)
    exact = history[history["Date"] == analog_date]
    values = pd.to_numeric(exact[column], errors="coerce").dropna()
    if not values.empty:
        return float(values.median())

    nearby = history[
        (history["Date"].dt.year == year)
        & ((history["Date"].dt.dayofyear - analog_date.dayofyear).abs() <= 3)
    ]
    values = pd.to_numeric(nearby[column], errors="coerce").dropna()
    return float(values.median()) if not values.empty else None


def actual_last_date(history: pd.DataFrame, column: str) -> pd.Timestamp:
    values = history[history[column].notna()] if column in history.columns else pd.DataFrame()
    if values.empty:
        return history["Date"].max()
    return pd.to_datetime(values["Date"]).max()


def analog_year_scores(history: pd.DataFrame, as_of_date: pd.Timestamp, column: str = "Q_in_m3s") -> list[tuple[int, float]]:
    current = history[
        (history["Date"].dt.year == FORECAST_YEAR)
        & (history["Date"] <= as_of_date)
        & history[column].notna()
    ][["Date", column]].copy()
    if current.empty:
        return []
    current["DayOfYear"] = current["Date"].dt.dayofyear
    current_values = pd.to_numeric(current[column], errors="coerce")

    scores: list[tuple[int, float]] = []
    for year in sorted(history["Date"].dt.year.dropna().astype(int).unique()):
        if year >= FORECAST_YEAR:
            continue
        candidate = history[
            (history["Date"].dt.year == year)
            & (history["Date"].dt.dayofyear.isin(current["DayOfYear"]))
            & history[column].notna()
        ][["Date", column]].copy()
        if len(candidate) < max(90, len(current) * 0.5):
            continue
        candidate["DayOfYear"] = candidate["Date"].dt.dayofyear
        merged = current[["DayOfYear", column]].merge(
            candidate[["DayOfYear", column]],
            on="DayOfYear",
            suffixes=("_current", "_candidate"),
        )
        if len(merged) < 90:
            continue
        current_series = pd.to_numeric(merged[f"{column}_current"], errors="coerce")
        candidate_series = pd.to_numeric(merged[f"{column}_candidate"], errors="coerce")
        diff = current_series - candidate_series
        rmse = float((diff.pow(2).mean()) ** 0.5)
        current_mean = float(current_values.mean())
        candidate_mean = float(candidate_series.mean())
        mean_penalty = abs(current_mean - candidate_mean)
        score = rmse + mean_penalty * 0.4

        for helper_column, helper_weight in [("Rain_Basin_mm", 0.18), ("WaterLevel_m", 0.12)]:
            if helper_column not in history.columns:
                continue
            helper_current = history[
                (history["Date"].dt.year == FORECAST_YEAR)
                & (history["Date"] <= as_of_date)
                & history[helper_column].notna()
            ][["Date", helper_column]].copy()
            if len(helper_current) < 90:
                continue
            helper_current["DayOfYear"] = helper_current["Date"].dt.dayofyear
            helper_candidate = history[
                (history["Date"].dt.year == year)
                & (history["Date"].dt.dayofyear.isin(helper_current["DayOfYear"]))
                & history[helper_column].notna()
            ][["Date", helper_column]].copy()
            helper_candidate["DayOfYear"] = helper_candidate["Date"].dt.dayofyear
            helper_merged = helper_current[["DayOfYear", helper_column]].merge(
                helper_candidate[["DayOfYear", helper_column]],
                on="DayOfYear",
                suffixes=("_current", "_candidate"),
            )
            if len(helper_merged) < 90:
                continue
            helper_diff = pd.to_numeric(helper_merged[f"{helper_column}_current"], errors="coerce") - pd.to_numeric(
                helper_merged[f"{helper_column}_candidate"],
                errors="coerce",
            )
            helper_rmse = float((helper_diff.pow(2).mean()) ** 0.5)
            score += helper_rmse * helper_weight
        scores.append((year, score))
    return sorted(scores, key=lambda item: item[1])


def multi_analog_value(
    history: pd.DataFrame,
    target_date: pd.Timestamp,
    analog_years: list[int],
    column: str,
    quantile: float = 0.65,
) -> float | None:
    values = [
        value
        for year in analog_years
        for value in [historical_value_for_date(history, year, target_date, column)]
        if value is not None and np.isfinite(value)
    ]
    if not values:
        return same_month_day(history, target_date, column)
    analog_value = float(pd.Series(values).quantile(quantile))
    climatology = same_month_day(history, target_date, column)
    if climatology is None:
        return analog_value
    return float(0.75 * analog_value + 0.25 * climatology)


def historical_window_value(
    history: pd.DataFrame,
    year: int,
    target_date: pd.Timestamp,
    column: str,
    radius_days: int = 3,
    reducer: str = "max",
) -> float | None:
    try:
        analog_date = target_date.replace(year=year)
    except ValueError:
        analog_date = target_date.replace(year=year, day=28)
    window = history[
        (history["Date"] >= analog_date - pd.Timedelta(days=radius_days))
        & (history["Date"] <= analog_date + pd.Timedelta(days=radius_days))
    ]
    values = pd.to_numeric(window[column], errors="coerce").dropna()
    if values.empty:
        return None
    if reducer == "median":
        return float(values.median())
    return float(values.max())


def storm_analog_value(history: pd.DataFrame, target_date: pd.Timestamp, analog_years: list[int], column: str) -> float | None:
    exact = multi_analog_value(
        history,
        target_date,
        analog_years,
        column,
        quantile=0.85 if target_date.month in STORM_MONTHS else 0.75,
    )
    if target_date.month not in STORM_MONTHS | RAINY_SHOULDER_MONTHS:
        return exact
    event_values = [
        value
        for year in analog_years
        for value in [historical_window_value(history, year, target_date, column, radius_days=4, reducer="max")]
        if value is not None and np.isfinite(value)
    ]
    if not event_values:
        return exact
    event_value = float(pd.Series(event_values).quantile(0.72 if target_date.month in STORM_MONTHS else 0.62))
    if exact is None:
        return event_value
    event_weight = 0.38 if target_date.month in STORM_MONTHS else 0.24
    return float((1.0 - event_weight) * exact + event_weight * event_value)


def q_blend_weight(date: pd.Timestamp) -> float:
    if date.month in STORM_MONTHS:
        return 0.25
    if date.month in RAINY_SHOULDER_MONTHS:
        return 0.35
    if date.month == 7:
        return 0.45
    return 0.62


def water_blend_weight(date: pd.Timestamp) -> float:
    if date.month in STORM_MONTHS:
        return 0.35
    if date.month in RAINY_SHOULDER_MONTHS:
        return 0.45
    if date.month == 7:
        return 0.55
    return 0.70


def recent_mean(history: pd.DataFrame, end_date: pd.Timestamp, column: str, days: int = 30) -> float | None:
    start_date = end_date - pd.Timedelta(days=days - 1)
    values = pd.to_numeric(
        history[(history["Date"] >= start_date) & (history["Date"] <= end_date)][column],
        errors="coerce",
    ).dropna()
    return float(values.mean()) if not values.empty else None


def analog_recent_mean(history: pd.DataFrame, end_date: pd.Timestamp, column: str, days: int = 30) -> float | None:
    analog_end = end_date.replace(year=ANALOG_YEAR)
    return recent_mean(history, analog_end, column, days=days)


def analog_recent_mean_multi(history: pd.DataFrame, end_date: pd.Timestamp, analog_years: list[int], column: str, days: int = 30) -> float | None:
    values = []
    for year in analog_years:
        try:
            analog_end = end_date.replace(year=year)
        except ValueError:
            analog_end = end_date.replace(year=year, day=28)
        value = recent_mean(history, analog_end, column, days=days)
        if value is not None:
            values.append(value)
    return float(pd.Series(values).median()) if values else None


def smooth(values: pd.Series, window: int) -> pd.Series:
    return values.rolling(window, center=True, min_periods=1).mean()


def clamp_water(basin_id: str, values: pd.Series) -> pd.Series:
    low, high = WATER_LEVEL_BOUNDS[basin_id]
    return values.clip(lower=low, upper=high)


def month_day(year: int, value: str) -> pd.Timestamp:
    month, day = [int(part) for part in value.split("-", 1)]
    return pd.Timestamp(year=year, month=month, day=day)


def apply_operation_water_rule(
    basin_id: str,
    water_values: pd.Series,
    cutoff: pd.Timestamp,
) -> tuple[pd.Series, bool]:
    """Apply reservoir operation lower-bound rules for seasonal filling."""
    rule = RESERVOIR_OPERATION_RULES.get(basin_id)
    if not rule:
        return water_values, False

    adjusted = water_values.copy()
    year = int(adjusted.index.max().year)
    ramp_start = month_day(year, str(rule["Ramp_Start"]))
    end_nov = pd.Timestamp(year=year, month=11, day=30)
    end_year = pd.Timestamp(year=year, month=12, day=31)
    max_level = float(rule["Max_Level_m"])
    end_nov_target = float(rule["End_Nov_Target_m"])
    end_year_target = float(rule["End_Year_Target_m"])

    start_value = adjusted.loc[adjusted.index <= ramp_start].dropna()
    if start_value.empty:
        start_level = float(adjusted.dropna().iloc[0]) if adjusted.notna().any() else end_nov_target
    else:
        start_level = float(start_value.iloc[-1])
    start_level = min(start_level, end_nov_target)

    changed = False
    for date in adjusted.index:
        if date <= cutoff or date < ramp_start:
            continue
        if date <= end_nov:
            progress = (date - ramp_start).days / max((end_nov - ramp_start).days, 1)
            target = start_level + progress * (end_nov_target - start_level)
        elif date <= end_year:
            progress = (date - end_nov).days / max((end_year - end_nov).days, 1)
            target = end_nov_target + progress * (end_year_target - end_nov_target)
        else:
            target = end_year_target
        target = min(target, max_level)
        if pd.isna(adjusted.loc[date]) or adjusted.loc[date] < target:
            adjusted.loc[date] = target
            changed = True
    return clamp_water(basin_id, adjusted.clip(upper=max_level)), changed


def recent_hydrology_momentum(history: pd.DataFrame, last_date: pd.Timestamp, column: str) -> float:
    observed = history[(history["Date"] <= last_date) & history[column].notna()][["Date", column]].tail(7).copy()
    if len(observed) < 3:
        return 0.0
    values = pd.to_numeric(observed[column], errors="coerce").dropna()
    if len(values) < 3:
        return 0.0
    last_value = float(values.iloc[-1])
    previous_mean = float(values.iloc[-4:-1].mean()) if len(values) >= 4 else float(values.iloc[:-1].mean())
    if not np.isfinite(previous_mean) or previous_mean == 0:
        return 0.0
    relative = (last_value - previous_mean) / abs(previous_mean)
    return float(np.clip(relative, -0.20, 0.75))


def apply_current_weather_boost(
    basin_id: str,
    history: pd.DataFrame,
    dates: pd.DatetimeIndex,
    q_values: pd.Series,
    water_values: pd.Series,
    last_q_date: pd.Timestamp,
    last_water_date: pd.Timestamp,
    weather: dict[str, object],
) -> tuple[pd.Series, pd.Series, bool]:
    """Boost short-term forecast using current wet weather and recent rising hydro trend."""
    wetness = basin_weather_wetness(basin_id, weather)
    q_momentum = recent_hydrology_momentum(history, last_q_date, "Q_in_m3s")
    water_momentum = recent_hydrology_momentum(history, last_water_date, "WaterLevel_m")
    if wetness <= 0.05 and q_momentum <= 0.05 and water_momentum <= 0.02:
        return q_values, water_values, False

    adjusted_q = q_values.copy()
    adjusted_water = water_values.copy()
    last_q_series = pd.to_numeric(history.loc[history["Date"] <= last_q_date, "Q_in_m3s"], errors="coerce").dropna()
    last_water_series = pd.to_numeric(history.loc[history["Date"] <= last_water_date, "WaterLevel_m"], errors="coerce").dropna()
    last_q = float(last_q_series.iloc[-1]) if not last_q_series.empty else None
    last_water = float(last_water_series.iloc[-1]) if not last_water_series.empty else None
    q_floor_factor = 1.0 + max(0.0, q_momentum) * 0.35 + wetness * 0.18
    q_multiplier = 1.0 + max(0.0, q_momentum) * 0.30 + wetness * 0.22
    water_lift_max = max(0.0, water_momentum) * 0.65 + wetness * 0.35
    changed = False
    for date in dates:
        lead = (date - last_q_date).days
        if lead <= 0 or lead > WEATHER_IMPACT_DAYS:
            continue
        decay = (WEATHER_IMPACT_DAYS - lead + 1) / WEATHER_IMPACT_DAYS
        if last_q is not None and pd.notna(adjusted_q.loc[date]):
            floor = last_q * (1.0 + (q_floor_factor - 1.0) * decay)
            boosted = adjusted_q.loc[date] * (1.0 + (q_multiplier - 1.0) * decay)
            new_q = max(float(adjusted_q.loc[date]), min(max(floor, boosted), last_q * 2.25))
            if new_q > adjusted_q.loc[date]:
                adjusted_q.loc[date] = new_q
                changed = True
        if last_water is not None and pd.notna(adjusted_water.loc[date]):
            water_target = last_water + water_lift_max * decay
            if adjusted_water.loc[date] < water_target:
                adjusted_water.loc[date] = water_target
                changed = True
    return adjusted_q.clip(lower=0), clamp_water(basin_id, adjusted_water), changed


def apply_kttv_forecast_adjustment(
    basin_id: str,
    dates: pd.DatetimeIndex,
    q_values: pd.Series,
    water_values: pd.Series,
    actual_cutoff: pd.Timestamp,
) -> tuple[pd.Series, pd.Series, bool]:
    """Blend official KTTV inflow forecasts into future Q and nudge water level consistently."""
    kttv = kttv_series_for_basin(basin_id, dates)
    if kttv.empty:
        return q_values, water_values, False

    adjusted_q = q_values.copy()
    original_q = q_values.copy()
    changed = False
    for row in kttv.to_dict("records"):
        date = pd.Timestamp(row["Date"]).normalize()
        if date <= actual_cutoff or date not in adjusted_q.index:
            continue
        kttv_q = to_float(row.get("KTTV_Q_m3s"))
        if kttv_q is None:
            continue
        source_type = str(row.get("Source_Type") or "")
        weight = KTTV_SOURCE_WEIGHTS.get(source_type, 0.35)
        current = adjusted_q.loc[date]
        if pd.isna(current):
            adjusted_q.loc[date] = max(0.0, kttv_q)
        else:
            adjusted_q.loc[date] = max(0.0, (1.0 - weight) * float(current) + weight * kttv_q)
        changed = True

    if not changed:
        return q_values, water_values, False

    q_delta = (adjusted_q - original_q).replace([np.inf, -np.inf], np.nan).fillna(0.0)
    water_response = KTTV_WATER_RESPONSE.get(basin_id, 0.05)
    water_adjustment = q_delta.rolling(5, min_periods=1).mean().clip(lower=-200.0, upper=500.0) / 100.0
    adjusted_water = water_values.copy()
    future_mask = adjusted_water.index > actual_cutoff
    adjusted_water.loc[future_mask] = adjusted_water.loc[future_mask] + water_adjustment.loc[future_mask] * water_response
    return adjusted_q.clip(lower=0), clamp_water(basin_id, adjusted_water), True


def apply_enso_hydrology_scenario(
    basin_id: str,
    history: pd.DataFrame,
    dates: pd.DatetimeIndex,
    q_values: pd.Series,
    water_values: pd.Series,
    actual_cutoff: pd.Timestamp,
) -> tuple[pd.Series, pd.Series, pd.Series, bool]:
    """Align monthly Q trend with the official ENSO/KTTV outlook without flattening daily events."""
    anomaly = pd.Series(
        [anomaly_for_date(basin_id, date) for date in dates],
        index=dates,
        dtype="float64",
    )
    future_scenario = (dates > actual_cutoff) & anomaly.notna().to_numpy()
    if not future_scenario.any():
        return q_values, water_values, anomaly, False

    original_q = pd.to_numeric(q_values, errors="coerce").copy()
    adjusted_q = original_q.copy()
    scale = pd.Series(1.0, index=dates, dtype="float64")
    historical = history[
        (history["Date"] < pd.Timestamp(f"{FORECAST_YEAR}-01-01"))
        & history["Q_in_m3s"].notna()
    ].copy()
    historical["Q_in_m3s"] = pd.to_numeric(historical["Q_in_m3s"], errors="coerce")
    historical["Month"] = historical["Date"].dt.month
    climatology = historical.groupby([historical["Date"].dt.year, "Month"])["Q_in_m3s"].mean()
    climatology = climatology.groupby("Month").median()

    for period in sorted(pd.PeriodIndex(dates[future_scenario], freq="M").unique()):
        month_mask = (dates.to_period("M") == period) & (dates > actual_cutoff)
        month_anomaly = anomaly.loc[month_mask].dropna()
        current = original_q.loc[month_mask].dropna()
        normal_q = climatology.get(period.month, np.nan)
        if month_anomaly.empty or current.empty or not np.isfinite(normal_q) or normal_q <= 0:
            continue
        target_mean = float(normal_q) * (1.0 + float(month_anomaly.iloc[0]) / 100.0)
        raw_factor = float(np.clip(target_mean / max(float(current.mean()), 0.01), 0.35, 1.35))
        # The seasonal bulletin constrains the monthly mean; daily peaks still come
        # from the AI, rainfall, storm and analog-year components.
        scenario_weight = 1.00
        scale.loc[month_mask] = 1.0 + scenario_weight * (raw_factor - 1.0)

    # Smooth transitions around month boundaries while preserving day-to-day shape.
    scale = scale.rolling(7, center=True, min_periods=1).mean().clip(0.40, 1.25)
    scale.loc[dates <= actual_cutoff] = 1.0
    adjusted_q = (original_q * scale).clip(lower=0)

    # Water level is not multiplied by the ENSO percentage. Instead, slower inflow
    # slows positive refill and slightly accelerates drawdown; reservoir operation
    # rules are applied afterwards by the main forecast pipeline.
    original_water = pd.to_numeric(water_values, errors="coerce").copy()
    adjusted_water = original_water.copy()
    daily_change = original_water.diff().fillna(0.0)
    q_ratio = (adjusted_q / original_q.replace(0, np.nan)).replace([np.inf, -np.inf], np.nan).fillna(1.0)
    response = pd.Series(1.0, index=dates, dtype="float64")
    rising = daily_change >= 0
    response.loc[rising] = 1.0 + 0.65 * (q_ratio.loc[rising] - 1.0)
    response.loc[~rising] = 1.0 + 0.25 * (1.0 - q_ratio.loc[~rising])
    response = response.clip(0.55, 1.35)

    first_future = dates[dates > actual_cutoff]
    if len(first_future):
        first_date = first_future[0]
        previous_dates = dates[dates < first_date]
        if len(previous_dates) and pd.notna(original_water.loc[previous_dates[-1]]):
            level = float(original_water.loc[previous_dates[-1]])
            for date in first_future:
                level += float(daily_change.loc[date]) * float(response.loc[date])
                adjusted_water.loc[date] = 0.65 * level + 0.35 * float(original_water.loc[date])

    return adjusted_q, clamp_water(basin_id, adjusted_water), anomaly, True


def forecast_dates() -> pd.DatetimeIndex:
    return pd.date_range(FORECAST_START_DATE, FORECAST_END_DATE, freq="D")


def load_model_bundle(basin_id: str, slug: str):
    candidates = sorted(
        MODEL_DIR.glob(f"model_{basin_id}_{slug}*.joblib"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    model_file = candidates[0] if candidates else MODEL_DIR / f"model_{basin_id}_{slug}.joblib"
    if not model_file.exists():
        return None
    return joblib.load(model_file)


def predict_model_values(history: pd.DataFrame, dates: pd.DatetimeIndex, basin_id: str, slug: str) -> pd.Series:
    bundle = load_model_bundle(basin_id, slug)
    if bundle is None:
        return pd.Series(np.nan, index=dates)

    features = bundle.get("features", [])
    available = history[history["Date"].isin(dates)].copy()
    if available.empty:
        return pd.Series(np.nan, index=dates)

    for column in features:
        if column not in available.columns:
            available[column] = np.nan
    prediction = bundle["model"].predict(available[features])
    values = pd.Series(prediction, index=pd.to_datetime(available["Date"]))
    if slug == "WaterLevel":
        anchor = pd.Series(np.nan, index=available.index, dtype="float64")
        for column in ["WaterLevel_Lag1", "WaterLevel_3d_mean", "WaterLevel_7d_mean", "WaterLevel_Lag2", "WaterLevel_Lag3"]:
            if column in available.columns:
                anchor = anchor.combine_first(pd.to_numeric(available[column], errors="coerce"))
        anchor.index = pd.to_datetime(available["Date"])
        anchor = anchor.reindex(values.index)
        max_daily_drift = MAX_DAILY_WATER_CHANGE.get(basin_id, 1.0)
        delta = values - anchor
        drift_mask = anchor.notna() & delta.abs().gt(max_daily_drift)
        values.loc[drift_mask] = anchor.loc[drift_mask] + delta.loc[drift_mask].clip(
            lower=-max_daily_drift,
            upper=max_daily_drift,
        )
    return values.reindex(dates)


def forecast_one(basin_id: str) -> pd.DataFrame:
    feature_file = OUTPUT_DIR / f"HydroAI_{basin_id}_features.csv"
    if not feature_file.exists():
        raise FileNotFoundError(f"Thieu {feature_file}. Hay chay 06_feature_engineering.py truoc.")

    history = pd.read_csv(feature_file, parse_dates=["Date"]).sort_values("Date")
    history = augment_history_with_dashboard_actuals(history, basin_id)
    dates = forecast_dates()
    weather = load_current_weather_impact()
    last_q_date = actual_last_date(history, "Q_in_m3s")
    last_water_date = actual_last_date(history, "WaterLevel_m")
    actual_cutoff = min(last_q_date, last_water_date)
    analog_scores = analog_year_scores(history, min(last_q_date, pd.Timestamp(f"{FORECAST_YEAR}-07-24")), "Q_in_m3s")
    analog_years = [year for year, _score in analog_scores[:ANALOG_TOP_N]]
    if not analog_years:
        analog_years = [year for year in range(2016, 2026) if year != FORECAST_YEAR]

    q_recent = recent_mean(history, last_q_date, "Q_in_m3s", days=30)
    q_analog_recent = analog_recent_mean_multi(history, last_q_date, analog_years, "Q_in_m3s", days=30)
    q_scale = 1.0
    if q_recent is not None and q_analog_recent and q_analog_recent > 0:
        q_scale = float(np.clip(q_recent / q_analog_recent, 0.75, 1.70))

    last_water = pd.to_numeric(history.loc[history["Date"] <= last_water_date, "WaterLevel_m"], errors="coerce").dropna()
    last_water_value = float(last_water.iloc[-1]) if not last_water.empty else None
    analog_anchor = same_month_day(history, last_water_date, "WaterLevel_m")
    water_offset = (last_water_value - analog_anchor) if last_water_value is not None and analog_anchor is not None else 0.0

    rows = []
    for date in dates:
        analog_q = storm_analog_value(history, date, analog_years, "Q_in_m3s")
        analog_water = storm_analog_value(history, date, analog_years, "WaterLevel_m")
        analog_rain = same_month_day(history, date, "Rain_Basin_mm")
        rows.append(
            {
                "Date": date,
                "Basin_ID": basin_id,
                "Forecast_Type": f"future_multi_analog_{'_'.join(str(year) for year in analog_years[:3])}",
                "Rain_Basin_mm": 0.0 if analog_rain is None else analog_rain,
                "Q_P50_m3s": np.nan if analog_q is None else max(0.0, analog_q * q_scale),
                "WaterLevel_P50_m": np.nan if analog_water is None else analog_water + water_offset,
            }
        )

    forecast = pd.DataFrame(rows)
    analog_q_series = pd.to_numeric(pd.Series([row["Q_P50_m3s"] for row in rows], index=dates), errors="coerce").clip(lower=0)
    analog_water_series = clamp_water(
        basin_id,
        smooth(pd.to_numeric(pd.Series([row["WaterLevel_P50_m"] for row in rows], index=dates), errors="coerce"), 5),
    )
    q_model = predict_model_values(history, dates, basin_id, "Q")
    water_model = predict_model_values(history, dates, basin_id, "WaterLevel")
    q_model = q_model.clip(lower=0)
    water_model = clamp_water(basin_id, water_model)

    q_actual = history.set_index("Date")["Q_in_m3s"].reindex(dates) if "Q_in_m3s" in history.columns else pd.Series(np.nan, index=dates)
    water_actual = history.set_index("Date")["WaterLevel_m"].reindex(dates) if "WaterLevel_m" in history.columns else pd.Series(np.nan, index=dates)
    q_anchor = pd.to_numeric(q_actual, errors="coerce").interpolate(method="time", limit_direction="both")
    water_anchor = pd.to_numeric(water_actual, errors="coerce").interpolate(method="time", limit_direction="both")
    q_blended = q_model.copy()
    water_blended = water_model.copy()
    for date in dates:
        if date <= last_q_date and pd.notna(q_anchor.loc[date]):
            q_blended.loc[date] = q_anchor.loc[date]
        else:
            analog_value = analog_q_series.loc[date]
            model_value = q_model.loc[date]
            if pd.notna(analog_value) and pd.notna(model_value):
                model_weight = q_blend_weight(date)
                q_blended.loc[date] = model_weight * model_value + (1.0 - model_weight) * analog_value
            elif pd.notna(analog_value):
                q_blended.loc[date] = analog_value
        if date <= last_water_date and pd.notna(water_anchor.loc[date]):
            water_blended.loc[date] = water_anchor.loc[date]
        else:
            analog_water_value = analog_water_series.loc[date]
            model_water_value = water_model.loc[date]
            if pd.notna(analog_water_value) and pd.notna(model_water_value):
                model_weight = water_blend_weight(date)
                water_blended.loc[date] = model_weight * model_water_value + (1.0 - model_weight) * analog_water_value
            elif pd.notna(analog_water_value):
                water_blended.loc[date] = analog_water_value

    q_cap = pd.to_numeric(history.loc[history["Date"].dt.month.isin([8, 9, 10, 11, 12]), "Q_in_m3s"], errors="coerce").quantile(0.985)
    if pd.notna(q_cap):
        q_blended = q_blended.clip(lower=0, upper=float(q_cap) * 1.08)
    q_blended, water_blended, weather_boost_applied = apply_current_weather_boost(
        basin_id,
        history,
        dates,
        q_blended,
        water_blended,
        last_q_date,
        last_water_date,
        weather,
    )
    q_blended, water_blended, kttv_adjustment_applied = apply_kttv_forecast_adjustment(
        basin_id,
        dates,
        q_blended,
        water_blended,
        actual_cutoff,
    )
    q_blended, water_blended, enso_anomaly, enso_adjustment_applied = apply_enso_hydrology_scenario(
        basin_id,
        history,
        dates,
        q_blended,
        water_blended,
        actual_cutoff,
    )
    water_blended, operation_rule_applied = apply_operation_water_rule(basin_id, water_blended, actual_cutoff)

    forecast["Q_P50_m3s"] = q_blended.values
    forecast["WaterLevel_P50_m"] = clamp_water(basin_id, water_blended).values
    forecast["Forecast_Type"] = np.where(
        pd.to_datetime(forecast["Date"]) <= actual_cutoff,
        "observed_anchor",
        rows[0]["Forecast_Type"] if rows else "future_multi_analog",
    )
    forecast.loc[pd.to_datetime(forecast["Date"]) > actual_cutoff, "Forecast_Type"] = (
        f"ai_blend_storm_analog_{'_'.join(str(year) for year in analog_years[:3])}"
    )
    if operation_rule_applied:
        future_mask = pd.to_datetime(forecast["Date"]) > actual_cutoff
        forecast.loc[future_mask, "Forecast_Type"] = forecast.loc[future_mask, "Forecast_Type"] + "_operation_rule"
    if weather_boost_applied:
        future_mask = pd.to_datetime(forecast["Date"]) > actual_cutoff
        forecast.loc[future_mask, "Forecast_Type"] = forecast.loc[future_mask, "Forecast_Type"] + "_weather_current_boost"
    if kttv_adjustment_applied:
        future_mask = pd.to_datetime(forecast["Date"]) > actual_cutoff
        forecast.loc[future_mask, "Forecast_Type"] = forecast.loc[future_mask, "Forecast_Type"] + "_kttv_blend"
    if enso_adjustment_applied:
        future_mask = (pd.to_datetime(forecast["Date"]) > actual_cutoff) & enso_anomaly.notna().to_numpy()
        forecast.loc[future_mask, "Forecast_Type"] = forecast.loc[future_mask, "Forecast_Type"] + "_enso_scenario"

    forecast["ENSO_Q_Anomaly_pct"] = enso_anomaly.values
    forecast["ENSO_Scenario_Source"] = np.where(enso_anomaly.notna(), ENSO_SOURCE_NAME, "")

    analog_mask = forecast["Q_P50_m3s"].isna()
    if analog_mask.any():
        forecast.loc[analog_mask, "Q_P50_m3s"] = analog_q_series.loc[analog_mask].values

    analog_water_mask = forecast["WaterLevel_P50_m"].isna()
    if analog_water_mask.any():
        forecast.loc[analog_water_mask, "WaterLevel_P50_m"] = analog_water_series.loc[analog_water_mask].values

    q_center = pd.Series(forecast["Q_P50_m3s"].values, index=pd.to_datetime(forecast["Date"]))
    water_center = pd.Series(forecast["WaterLevel_P50_m"].values, index=pd.to_datetime(forecast["Date"]))
    q_p10, q_p90 = conformal_bounds(q_center, read_conformal_residuals(basin_id, "Q"), floor=0.0)
    water_p10, water_p90 = conformal_bounds(water_center, read_conformal_residuals(basin_id, "WaterLevel"))
    forecast["Q_P10_m3s"] = q_p10.values
    forecast["Q_P90_m3s"] = q_p90.values
    forecast["WaterLevel_P10_m"] = clamp_water(basin_id, water_p10).values
    forecast["WaterLevel_P90_m"] = clamp_water(basin_id, water_p90).values
    forecast["Q_P10_m3s"] = np.minimum(forecast["Q_P10_m3s"], forecast["Q_P50_m3s"])
    forecast["Q_P90_m3s"] = np.maximum(forecast["Q_P90_m3s"], forecast["Q_P50_m3s"])
    forecast["WaterLevel_P10_m"] = np.minimum(forecast["WaterLevel_P10_m"], forecast["WaterLevel_P50_m"])
    forecast["WaterLevel_P90_m"] = np.maximum(forecast["WaterLevel_P90_m"], forecast["WaterLevel_P50_m"])
    rule = RESERVOIR_OPERATION_RULES.get(basin_id)
    if rule:
        max_level = float(rule["Max_Level_m"])
        for column in ["WaterLevel_P10_m", "WaterLevel_P50_m", "WaterLevel_P90_m"]:
            forecast[column] = forecast[column].clip(upper=max_level)

    forecast = forecast[
        [
            "Date",
            "Basin_ID",
            "Forecast_Type",
            "ENSO_Q_Anomaly_pct",
            "ENSO_Scenario_Source",
            "Rain_Basin_mm",
            "Q_P10_m3s",
            "Q_P50_m3s",
            "Q_P90_m3s",
            "WaterLevel_P10_m",
            "WaterLevel_P50_m",
            "WaterLevel_P90_m",
        ]
    ]
    forecast.to_csv(OUTPUT_DIR / f"Forecast_{basin_id}.csv", index=False, encoding="utf-8-sig")
    print(
        f"Da du bao 2026 {basin_id}: {len(forecast)} ngay; "
        f"actual_cutoff={actual_cutoff.date()}; scale Q={q_scale:.2f}; analog years={analog_years}."
    )
    return forecast


def forecast_all() -> pd.DataFrame:
    ensure_project_dirs()
    enso_scenario = save_scenario_to_database(HOCHUA_DB_PATH)
    print(f"Da luu {len(enso_scenario)} dong kich ban ENSO/KTTV vao database.")
    refresh_kttv_forecast_database()
    frames: list[pd.DataFrame] = []
    for basin_id in BASIN_NAMES:
        try:
            frames.append(forecast_one(basin_id))
        except Exception as exc:
            print(f"Khong du bao {basin_id}: {exc}")
    if not frames:
        raise RuntimeError("Khong du bao duoc luu vuc nao.")

    result = pd.concat(frames, ignore_index=True)
    try:
        result.to_excel(FORECAST_FILE, index=False)
        print(f"Da xuat: {FORECAST_FILE}")
    except PermissionError:
        fallback_file = FORECAST_FILE.with_name(f"{FORECAST_FILE.stem}_new{FORECAST_FILE.suffix}")
        result.to_excel(fallback_file, index=False)
        print(f"File dang bi khoa, da xuat ban moi: {fallback_file}")

    sqlite_df = result.copy()
    sqlite_df["Date"] = pd.to_datetime(sqlite_df["Date"], errors="coerce").dt.strftime("%Y-%m-%d")
    with sqlite3.connect(HOCHUA_DB_PATH) as conn:
        sqlite_df.to_sql("hydroai_forecast_daily", conn, if_exists="replace", index=False)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_hydroai_forecast_daily_basin_date ON hydroai_forecast_daily(Basin_ID, Date)")
    print(f"Da luu du bao tuong lai vao database: {HOCHUA_DB_PATH}")
    write_forecast_report(result)
    return result


def write_forecast_report(forecast: pd.DataFrame) -> None:
    """Write daily forecast quality report with actuals and errors where available."""
    frames: list[pd.DataFrame] = []
    for basin_id in BASIN_NAMES:
        feature_file = OUTPUT_DIR / f"HydroAI_{basin_id}_features.csv"
        if not feature_file.exists():
            continue
        actual = pd.read_csv(
            feature_file,
            parse_dates=["Date"],
            usecols=lambda column: column in ["Date", "Basin_ID", "Q_in_m3s", "WaterLevel_m", "Rain_Basin_mm"],
            low_memory=False,
        )
        basin_forecast = forecast[forecast["Basin_ID"] == basin_id].copy()
        merged = basin_forecast.merge(actual, on=["Date", "Basin_ID"], how="left")
        merged["Q_Error_m3s"] = merged["Q_P50_m3s"] - merged["Q_in_m3s"]
        merged["WaterLevel_Error_m"] = merged["WaterLevel_P50_m"] - merged["WaterLevel_m"]
        merged["P10_P90_Covers_Q"] = (
            (merged["Q_in_m3s"] >= merged["Q_P10_m3s"])
            & (merged["Q_in_m3s"] <= merged["Q_P90_m3s"])
        ).where(merged["Q_in_m3s"].notna())
        merged["P10_P90_Covers_WaterLevel"] = (
            (merged["WaterLevel_m"] >= merged["WaterLevel_P10_m"])
            & (merged["WaterLevel_m"] <= merged["WaterLevel_P90_m"])
        ).where(merged["WaterLevel_m"].notna())
        frames.append(merged)
    if not frames:
        return
    report = pd.concat(frames, ignore_index=True)
    report["Year"] = report["Date"].dt.year
    report["Month"] = report["Date"].dt.month
    monthly = (
        report.groupby(["Basin_ID", "Year", "Month"], dropna=False)
        .agg(
            Actual_Q_Mean=("Q_in_m3s", "mean"),
            Forecast_Q_Mean=("Q_P50_m3s", "mean"),
            Q_MAE=("Q_Error_m3s", lambda values: values.abs().mean()),
            Actual_WaterLevel_Mean=("WaterLevel_m", "mean"),
            Forecast_WaterLevel_Mean=("WaterLevel_P50_m", "mean"),
            WaterLevel_MAE=("WaterLevel_Error_m", lambda values: values.abs().mean()),
            Q_Coverage=("P10_P90_Covers_Q", "mean"),
            WaterLevel_Coverage=("P10_P90_Covers_WaterLevel", "mean"),
        )
        .reset_index()
    )
    yearly = (
        report.groupby(["Basin_ID", "Year"], dropna=False)
        .agg(
            Q_MAE=("Q_Error_m3s", lambda values: values.abs().mean()),
            WaterLevel_MAE=("WaterLevel_Error_m", lambda values: values.abs().mean()),
            Q_Coverage=("P10_P90_Covers_Q", "mean"),
            WaterLevel_Coverage=("P10_P90_Covers_WaterLevel", "mean"),
        )
        .reset_index()
    )
    output_file = OUTPUT_DIR / "Forecast_Report.xlsx"
    try:
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            report.to_excel(writer, sheet_name="Daily", index=False)
            monthly.to_excel(writer, sheet_name="Monthly", index=False)
            yearly.to_excel(writer, sheet_name="Yearly", index=False)
        print(f"Da xuat bao cao du bao: {output_file}")
    except PermissionError:
        fallback_file = output_file.with_name(f"{output_file.stem}_new{output_file.suffix}")
        with pd.ExcelWriter(fallback_file, engine="openpyxl") as writer:
            report.to_excel(writer, sheet_name="Daily", index=False)
            monthly.to_excel(writer, sheet_name="Monthly", index=False)
            yearly.to_excel(writer, sheet_name="Yearly", index=False)
        print(f"File bao cao dang bi khoa, da xuat ban moi: {fallback_file}")


if __name__ == "__main__":
    try:
        forecast_all()
    except Exception as exc:
        print("\nCHUONG TRINH DUNG DO LOI:")
        print(exc)
        sys.exit(1)
